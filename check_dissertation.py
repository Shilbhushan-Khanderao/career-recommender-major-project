from docx import Document

doc = Document('docs/MTech_Major_Dissertation_Final.docx')

print("=" * 70)
print("TABLE VERIFICATION REPORT")
print("=" * 70)

# 1. Find all Table captions in paragraphs
print("\n### TABLE CAPTIONS IN DOCUMENT BODY ###\n")
table_captions = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text.startswith('Table ') and ':' in text:
        table_captions.append((i, text))
        print(f"[Para {i}] {text}")
    elif 'Table' in text and any(x in text for x in ['Table C.', 'Table E.', 'Table F.', 'Table 3.', 'Table 4.']):
        if ':' in text or text.startswith('Table'):
            print(f"[Para {i}] {text[:100]}")

# 2. Show all actual table objects with context
print("\n### ALL TABLE OBJECTS IN DOCUMENT ###\n")
for idx, table in enumerate(doc.tables):
    first_cell = table.rows[0].cells[0].text.strip() if table.rows else ''
    num_rows = len(table.rows)
    num_cols = len(table.rows[0].cells) if table.rows else 0
    print(f"Table {idx}: {num_rows} rows x {num_cols} cols | First cell: '{first_cell[:50]}'")

# 3. Check Appendix sections for table headings
print("\n### APPENDIX TABLE HEADINGS ###\n")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if i > 500:  # Focus on appendix area
        if ('C.' in text or 'E.' in text or 'F.' in text) and len(text) < 80:
            print(f"[Para {i}] {text}")

# 2. Find all figure captions in body
print("\n### FIGURE CAPTIONS IN DOCUMENT BODY ###\n")
figures_in_body = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text.startswith('Figure ') and ':' in text:
        figures_in_body.append(text)
        print(f"[Para {i}] {text}")

# 3. Find figure references (like "Figure E.1 (Appendix E)")
print("\n### FIGURE REFERENCES (in text) ###\n")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if 'Figure' in text and '(Appendix' in text:
        print(f"[Para {i}] {text[:100]}...")

# 4. Count images
print("\n### IMAGE COUNT ###")
image_count = 0
for rel in doc.part.rels.values():
    if 'image' in rel.reltype:
        image_count += 1
print(f"Total images in document: {image_count}")

# 5. Check Chapter structure
print("\n### CHAPTER STRUCTURE ###\n")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text.startswith('CHAPTER') or text.startswith('Chapter'):
        print(f"[Para {i}] {text}")
    if text.startswith('APPENDIX') or (text.startswith('Appendix') and len(text) < 50):
        print(f"[Para {i}] {text}")

print("\n" + "=" * 70)
print("VERIFICATION COMPLETE")
print("=" * 70)
