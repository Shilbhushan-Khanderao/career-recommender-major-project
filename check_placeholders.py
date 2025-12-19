from docx import Document

doc = Document('docs/MTech_Major_Dissertation_Final.docx')

print("=" * 70)
print("PLACEHOLDER & IRRELEVANT TEXT CHECK")
print("=" * 70)

# Common placeholder patterns
placeholders = ['TODO', 'FIXME', 'XXX', 'placeholder', 'Placeholder', 'PLACEHOLDER', 
                'insert', 'Insert', 'add here', 'Add here', 'to be added', 'TBD',
                '[...]', '...add', 'Lorem', 'ipsum', 'example text',
                'fill in', 'your name', 'YOUR NAME', 'delete this',
                'copy paste', 'copied from', 'remove this', 'update this',
                'change this', 'replace this', 'edit this', 'modify this']

print("\n### CHECKING FOR PLACEHOLDERS ###\n")
found_placeholders = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        for ph in placeholders:
            if ph.lower() in text.lower():
                found_placeholders.append((i, ph, text[:100]))
                print(f"[Para {i}] Found '{ph}': {text[:80]}...")

if not found_placeholders:
    print("No placeholder text found!")

# Check for blank spaces (underscores used as fill-in blanks)
print("\n### CHECKING FOR BLANK SPACES (___) ###\n")
blank_found = False
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if '___' in text:
        blank_found = True
        print(f"[Para {i}] {text[:100]}")

if not blank_found:
    print("No blank spaces found!")

# Check for incomplete sentences or fragments
print("\n### CHECKING FOR INCOMPLETE/SUSPICIOUS TEXT ###\n")
suspicious = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        # Check for text ending with incomplete patterns
        if text.endswith('...') and len(text) > 10:
            print(f"[Para {i}] Ends with '...': {text[:80]}")
            suspicious.append(i)
        # Check for standalone brackets
        if text == '[]' or text == '()' or text == '{}':
            print(f"[Para {i}] Empty brackets: {text}")
            suspicious.append(i)
        # Check for "Figure X" or "Table X" without proper caption
        if (text.startswith('Figure ') or text.startswith('Table ')) and ':' not in text and len(text) < 15:
            print(f"[Para {i}] Incomplete caption: {text}")
            suspicious.append(i)

# Check for copy-paste artifacts
print("\n### CHECKING FOR COPY-PASTE ARTIFACTS ###\n")
artifacts = ['http://', 'https://', 'www.', 'file:///', 'C:\\', 'D:\\', 
             'Error:', 'Warning:', 'Exception:', 'Traceback']
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        for artifact in artifacts:
            if artifact in text and i > 50:  # Skip front matter
                # Skip if it's in references section
                if i < 480 or i > 510:  # References are around para 480-510
                    print(f"[Para {i}] Found '{artifact}': {text[:80]}...")

# Check for duplicate paragraphs
print("\n### CHECKING FOR DUPLICATE PARAGRAPHS ###\n")
seen_texts = {}
duplicates = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text and len(text) > 50:  # Only check substantial paragraphs
        if text in seen_texts:
            duplicates.append((i, seen_texts[text], text[:60]))
        else:
            seen_texts[text] = i

if duplicates:
    for dup in duplicates[:10]:  # Show first 10
        print(f"[Para {dup[0]}] Duplicate of Para {dup[1]}: {dup[2]}...")
else:
    print("No duplicate paragraphs found!")

# Check for unusual characters
print("\n### CHECKING FOR UNUSUAL CHARACTERS ###\n")
unusual_chars = ['�', '□', '■', '▪', '►', '◄', '○', '●']
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        for char in unusual_chars:
            if char in text:
                print(f"[Para {i}] Found unusual char '{char}': {text[:60]}...")

# Check tables for placeholder content
print("\n### CHECKING TABLES FOR PLACEHOLDERS ###\n")
for idx, table in enumerate(doc.tables):
    for row_idx, row in enumerate(table.rows):
        for cell_idx, cell in enumerate(row.cells):
            cell_text = cell.text.strip()
            if any(ph.lower() in cell_text.lower() for ph in ['placeholder', 'tbd', 'todo', 'xxx', '___']):
                print(f"Table {idx}, Row {row_idx}, Cell {cell_idx}: {cell_text[:50]}")

print("\n" + "=" * 70)
print("CHECK COMPLETE")
print("=" * 70)
